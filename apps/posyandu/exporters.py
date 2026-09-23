"""Helper export Master Data Admin Kelurahan ke Excel dan PDF."""

from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _download_name(title: str, extension: str) -> str:
    stamp = timezone.localdate().strftime("%Y%m%d")
    return f"{slugify(title) or 'master-data'}_{stamp}.{extension}"


def master_excel_response(*, title: str, headers: list[str], rows: list[list], filter_text: str = ""):
    """Buat file XLSX rapi untuk master data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Data"

    total_columns = max(len(headers), 1)
    end_column = get_column_letter(total_columns)

    ws.merge_cells(f"A1:{end_column}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{end_column}2")
    metadata = f"Diekspor: {timezone.localtime().strftime('%d-%m-%Y %H:%M')}"
    if filter_text:
        metadata += f" | Filter: {filter_text}"
    ws["A2"] = metadata
    ws["A2"].font = Font(italic=True, size=10)
    ws["A2"].alignment = Alignment(horizontal="center")

    header_row = 4
    thin = Side(style="thin", color="D9E1F2")
    header_fill = PatternFill("solid", fgColor="DCE6F1")

    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{end_column}{max(header_row, header_row + len(rows))}"

    for column_index, header in enumerate(headers, start=1):
        values = [str(header)]
        values.extend(str(row[column_index - 1] if column_index - 1 < len(row) else "") for row in rows)
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), 45)
        ws.column_dimensions[get_column_letter(column_index)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type=EXCEL_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{_download_name(title, "xlsx")}"'
    return response


def master_pdf_response(*, title: str, headers: list[str], rows: list[list], filter_text: str = ""):
    """Buat PDF A4 landscape dengan tabel master data."""
    html = render_to_string(
        "posyandu/export/master_data_pdf.html",
        {
            "title": title,
            "headers": headers,
            "rows": rows,
            "filter_text": filter_text,
            "generated_at": timezone.localtime(),
            "total_data": len(rows),
        },
    )

    buffer = BytesIO()
    result = pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
    if result.err:
        return HttpResponse("Gagal membuat PDF master data.", status=500)

    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_download_name(title, "pdf")}"'
    return response
