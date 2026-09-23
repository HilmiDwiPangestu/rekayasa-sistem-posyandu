from datetime import date, time, datetime, timedelta

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.db.models import Count
from django.utils import timezone

from apps.accounts.models import Petugas
from apps.posyandu.forms import JadwalKegiatanForm
from apps.posyandu.models import JadwalKegiatan, Posyandu


class JadwalServiceTests(TestCase):
    def setUp(self):
        self.pos_a = Posyandu.objects.create(nama="Pos A")
        self.pos_b = Posyandu.objects.create(nama="Pos B")
        self.user = User.objects.create_user("kader-jadwal", password="aman-12345")
        Petugas.objects.create(
            user=self.user,
            nama="Kader Jadwal",
            level="kader",
            posyandu=self.pos_a,
        )
        self.jadwal_a = JadwalKegiatan.objects.create(
            tgl_kegiatan=date.today(),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_a,
        )
        JadwalKegiatan.objects.create(
            tgl_kegiatan=date.today(),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_b,
        )

    def test_kader_hanya_melihat_jadwal_posyandunya(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("posyandu:list_jadwal"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].paginator.count, 1)

    def test_kader_tidak_dapat_mengubah_jadwal_pos_lain(self):
        self.client.force_login(self.user)
        jadwal_lain = JadwalKegiatan.objects.get(posyandu=self.pos_b)
        response = self.client.get(reverse("posyandu:edit", args=[jadwal_lain.pk]))
        self.assertEqual(response.status_code, 404)

    def test_jam_selesai_harus_setelah_jam_mulai(self):
        form = JadwalKegiatanForm(data={
            "tgl_kegiatan": date.today(),
            "jam_mulai": "10:00",
            "jam_selesai": "09:00",
            "jns_kegiatan": "Imunisasi",
        })
        self.assertFalse(form.is_valid())
        self.assertIn("jam_selesai", form.errors)

    def test_form_jadwal_tidak_menampilkan_input_posyandu(self):
        form = JadwalKegiatanForm()
        self.assertNotIn("posyandu", form.fields)

    def test_create_jadwal_memakai_tempat_tugas_kader_dan_mengabaikan_post_posyandu(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("posyandu:create"),
            {
                "tgl_kegiatan": date.today(),
                "jam_mulai": "13:00",
                "jam_selesai": "15:00",
                "jns_kegiatan": "Pemeriksaan Rutin",
                # Upaya manipulasi lokasi harus diabaikan karena field bukan bagian form.
                "posyandu": self.pos_b.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        jadwal = JadwalKegiatan.objects.filter(jam_mulai=time(13)).latest("pk")
        self.assertEqual(jadwal.posyandu, self.pos_a)

    def test_breadcrumb_edit_memiliki_induk_jadwal(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("posyandu:edit", args=[self.jadwal_a.pk]))
        breadcrumb = response.context["django_breadcrumbs"]
        self.assertEqual(breadcrumb[0]["name"], "Jadwal Kegiatan")
        self.assertEqual(breadcrumb[-1]["name"], "Ubah Data")


class AdminPageSmokeTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("admin-uji", password="aman-12345")
        Petugas.objects.create(
            user=self.user,
            nama="Admin Uji",
            level="admin",
        )
        self.client.force_login(self.user)

    def test_seluruh_halaman_utama_admin_dapat_dibuka(self):
        names = [
            "dashboard:index",
            "posyandu:daftar_petugas_kader",
            "posyandu:daftar_bidan",
            "posyandu:list_posyandu",
            "posyandu:create_posyandu",
            "posyandu:jadwal_posyandu_admin",
            "posyandu:semua_laporan",
            "posyandu:laporan_bumil",
            "posyandu:laporan_balita",
            "posyandu:laporan_imunisasi",
            "posyandu:laporan_triwulan",
            "laporan:laporan_index",
            "laporan:laporan_bulanan",
            "laporan:laporan_triwulan",
        ]
        for name in names:
            with self.subTest(name=name):
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 200)


    def test_route_imunisasi_umum_admin_dialihkan_ke_laporan_admin(self):
        url_umum = reverse("laporan:laporan_imunisasi_bulanan")
        url_admin = reverse("posyandu:laporan_imunisasi")

        # URL harus berbeda agar redirect Admin tidak kembali ke route lama.
        self.assertNotEqual(url_umum, url_admin)
        self.assertEqual(url_umum, "/laporan/imunisasi/")
        self.assertEqual(url_admin, "/laporan/admin/imunisasi/")

        response = self.client.get(url_umum)
        self.assertRedirects(
            response,
            url_admin,
            fetch_redirect_response=False,
        )

        response_final = self.client.get(url_admin)
        self.assertEqual(response_final.status_code, 200)

    def test_ekspor_utama_menghasilkan_berkas(self):
        # Export imunisasi semua Posyandu membutuhkan minimal satu master Posyandu.
        Posyandu.objects.create(nama="Pos Smoke Export")
        exports = [
            ("posyandu:export_bumil_excel", "spreadsheetml"),
            ("posyandu:cetak_bumil_pdf", "application/pdf"),
            ("posyandu:cetak_balita_pdf", "application/pdf"),
            ("posyandu:cetak_imunisasi_pdf", "application/pdf"),
            ("laporan:export_laporan_triwulan_pdf", "application/pdf"),
        ]
        for name, expected_type in exports:
            with self.subTest(name=name):
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected_type, response["Content-Type"])
                self.assertGreater(len(response.content), 0)

    def test_semua_laporan_admin_menampilkan_empat_laporan_utama(self):
        response = self.client.get(reverse("posyandu:semua_laporan"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Laporan Pemeriksaan Balita")
        self.assertContains(response, "Laporan Pemeriksaan Ibu Hamil")
        self.assertContains(response, "Laporan Imunisasi")
        self.assertContains(response, "Laporan Triwulan")
        self.assertContains(response, reverse("posyandu:laporan_balita"))
        self.assertContains(response, reverse("posyandu:laporan_bumil"))
        self.assertContains(response, reverse("posyandu:laporan_imunisasi"))
        self.assertContains(response, reverse("posyandu:laporan_triwulan"))

    def test_laporan_triwulan_admin_menampilkan_opsi_pdf_terpilih_dan_semua(self):
        pos_a = Posyandu.objects.create(nama="Pos Admin PDF A")
        pos_b = Posyandu.objects.create(nama="Pos Admin PDF B")

        response = self.client.get(
            reverse("posyandu:laporan_triwulan"),
            {"tahun": 2026, "triwulan": 3, "posyandu": pos_b.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_posyandu_id"], pos_b.pk)
        self.assertContains(response, "Semua Posyandu · 1 PDF")
        self.assertContains(response, "scope=all")
        self.assertContains(response, "scope=selected")
        self.assertContains(response, "Semua Balita")
        self.assertContains(response, "Semua Ibu Hamil")
        self.assertEqual(
            [item.pk for item in response.context["daftar_posyandu_filter"]],
            [pos_a.pk, pos_b.pk],
        )


class DestructiveActionTests(TestCase):
    def setUp(self):
        self.pos = Posyandu.objects.create(nama="Pos Hapus")
        self.user = User.objects.create_user("kader-hapus", password="aman-12345")
        Petugas.objects.create(
            user=self.user,
            nama="Kader Hapus",
            level="kader",
            posyandu=self.pos,
        )
        self.jadwal = JadwalKegiatan.objects.create(
            tgl_kegiatan=date.today(),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos,
        )
        self.client.force_login(self.user)

    def test_hapus_jadwal_hanya_post(self):
        url = reverse("posyandu:delete", args=[self.jadwal.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url)
        self.assertEqual(response.status_code, 302)
        self.assertFalse(JadwalKegiatan.objects.filter(pk=self.jadwal.pk).exists())


class PosyanduAccessHelperTests(TestCase):
    def test_kader_hanya_memakai_posyandu_utama_sebagai_akses_operasional(self):
        from common.access import active_posyandu_ids
        from apps.posyandu.models import PenugasanPetugas

        utama = Posyandu.objects.create(nama="Pos Utama")
        aktif = Posyandu.objects.create(nama="Pos Tambahan Aktif")
        tambahan = Posyandu.objects.create(nama="Pos Tambahan Lain")
        user = User.objects.create_user("kader-akses", password="aman-12345")
        petugas = Petugas.objects.create(
            user=user,
            nama="Kader Akses",
            level="kader",
            posyandu=utama,
        )
        PenugasanPetugas.objects.create(
            petugas=petugas,
            posyandu=aktif,
            tanggal_daftar=date.today(),
        )
        PenugasanPetugas.objects.create(
            petugas=petugas,
            posyandu=tambahan,
            tanggal_daftar=date.today(),
        )

        self.assertEqual(active_posyandu_ids(petugas), {utama.pk})


class BidanCoverageServiceTests(TestCase):
    def setUp(self):
        from common.testing import make_petugas
        self.pos = [Posyandu.objects.create(nama=f"Pos {i}", desa="Kelurahan Uji") for i in range(1, 7)]
        self.user_a, self.bidan_a = make_petugas(
            username="bidan-cakupan-a",
            level="bidan",
            posyandu=self.pos[0],
        )
        self.user_b, self.bidan_b = make_petugas(
            username="bidan-cakupan-b",
            level="bidan",
            posyandu=self.pos[3],
        )

    def test_satu_bidan_dapat_memiliki_banyak_posyandu(self):
        from apps.posyandu.services import set_bidan_coverage, bidan_coverage_ids

        set_bidan_coverage(self.bidan_a, [self.pos[0].pk, self.pos[1].pk, self.pos[2].pk])
        self.assertEqual(
            bidan_coverage_ids(self.bidan_a),
            {self.pos[0].pk, self.pos[1].pk, self.pos[2].pk},
        )

    def test_cakupan_bidan_tidak_boleh_tumpang_tindih(self):
        from apps.posyandu.services import BidanCoverageError, set_bidan_coverage

        set_bidan_coverage(self.bidan_a, [self.pos[0].pk, self.pos[1].pk, self.pos[2].pk])
        with self.assertRaises(BidanCoverageError):
            set_bidan_coverage(self.bidan_b, [self.pos[2].pk, self.pos[3].pk])

    def test_mengubah_cakupan_menghapus_penugasan_lama(self):
        from apps.posyandu.services import set_bidan_coverage

        set_bidan_coverage(self.bidan_a, [self.pos[0].pk, self.pos[1].pk])
        set_bidan_coverage(self.bidan_a, [self.pos[1].pk, self.pos[2].pk])

        self.assertFalse(
            PenugasanPetugas.objects.filter(
                petugas=self.bidan_a,
                posyandu=self.pos[0],
            ).exists()
        )
        self.assertTrue(
            PenugasanPetugas.objects.filter(
                petugas=self.bidan_a,
                posyandu=self.pos[2],
            ).exists()
        )


class BidanCoverageViewTests(TestCase):
    def setUp(self):
        from common.testing import make_petugas
        self.admin, _ = make_petugas(username="admin-cakupan", level="admin", posyandu=None)
        self.pos1 = Posyandu.objects.create(nama="Pos 1", desa="Kelurahan Uji")
        self.pos2 = Posyandu.objects.create(nama="Pos 2", desa="Kelurahan Uji")
        _, self.bidan = make_petugas(
            username="bidan-cakupan-view",
            level="bidan",
            posyandu=self.pos1,
        )
        self.client.force_login(self.admin)

    def test_admin_dapat_membuka_halaman_cakupan_bidan(self):
        response = self.client.get(reverse("posyandu:atur_cakupan_bidan", args=[self.bidan.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Cakupan Wilayah Bidan")
        self.assertContains(response, self.pos1.nama)
        self.assertContains(response, self.pos2.nama)

    def test_admin_dapat_menyimpan_multi_posyandu(self):
        response = self.client.post(
            reverse("posyandu:atur_cakupan_bidan", args=[self.bidan.pk]),
            {"cakupan_posyandu": [self.pos1.pk, self.pos2.pk]},
        )
        self.assertEqual(response.status_code, 302)
        ids = set(
            PenugasanPetugas.objects.filter(
                petugas=self.bidan,
            ).values_list("posyandu_id", flat=True)
        )
        self.assertEqual(ids, {self.pos1.pk, self.pos2.pk})


class BidanCoveragePermissionTests(TestCase):
    def test_bidan_tidak_memiliki_akun_login_dan_cakupan_hanya_diatur_admin(self):
        from common.testing import make_petugas

        pos = Posyandu.objects.create(nama="Pos Permission")
        user, bidan = make_petugas(
            username="bidan-permission-cakupan",
            level="bidan",
            posyandu=pos,
        )
        self.assertIsNone(user)
        self.assertIsNone(bidan.user_id)

        admin, _ = make_petugas(username="admin-permission-cakupan", level="admin", posyandu=None)
        self.client.force_login(admin)
        response = self.client.get(reverse("posyandu:atur_cakupan_bidan", args=[bidan.pk]))
        self.assertEqual(response.status_code, 200)


class KaderAccountCreationViewTests(TestCase):
    def setUp(self):
        from common.testing import make_petugas, make_posyandu
        from apps.accounts.models import Petugas

        self.posyandu = make_posyandu("Pos Akun Kader")
        self.admin, _ = make_petugas(
            username="admin-akun-kader",
            level="admin",
            posyandu=None,
        )
        self.kader = Petugas.objects.create(
            nama="Kader Akun",
            level="kader",
            posyandu=self.posyandu,
            no_telp="081277788899",
        )
        self.client.force_login(self.admin)

    def test_form_manual_membuat_username_dan_password_sesuai_input(self):
        response = self.client.post(
            reverse("posyandu:buat_akun_petugas", args=[self.kader.pk]),
            {
                "username": "kader.custom",
                "password": "SandiKader!2026",
                "password_confirm": "SandiKader!2026",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.kader.refresh_from_db()
        self.assertEqual(self.kader.user.username, "kader.custom")
        self.assertTrue(self.kader.user.check_password("SandiKader!2026"))

    def test_form_kosong_tetap_generate_otomatis(self):
        response = self.client.post(
            reverse("posyandu:buat_akun_petugas", args=[self.kader.pk]),
            {"username": "", "password": "", "password_confirm": ""},
        )
        self.assertEqual(response.status_code, 302)
        self.kader.refresh_from_db()
        self.assertEqual(self.kader.user.username, "081277788899")
        self.assertTrue(self.kader.user.has_usable_password())

    def test_konfirmasi_password_tidak_sama_tidak_membuat_akun(self):
        response = self.client.post(
            reverse("posyandu:buat_akun_petugas", args=[self.kader.pk]),
            {
                "username": "kader.gagal",
                "password": "SandiKader!2026",
                "password_confirm": "BedaPassword!2026",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.kader.refresh_from_db()
        self.assertIsNone(self.kader.user_id)

    def test_daftar_kader_memuat_modal_dua_opsi_akun(self):
        response = self.client.get(reverse("posyandu:daftar_petugas_kader"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Pilih cara membuat akun")
        self.assertContains(response, "Otomatis")
        self.assertContains(response, "Atur Sendiri")
        self.assertContains(response, 'id="accountAutoPanel"')
        self.assertContains(response, 'id="accountManualPanel"')
        self.assertContains(response, 'id="accountUsername"')
        self.assertContains(response, 'id="accountPassword"')
        self.assertContains(response, 'id="accountSubmitLabel"')

class UnifiedSeedCommandTests(TestCase):
    def test_seed_data_menyiapkan_seluruh_data_demo_dengan_satu_command(self):
        from io import StringIO
        from django.core.management import call_command
        from apps.peserta.models import Peserta
        from apps.referensi.models import WHOStandardPBU
        from apps.posyandu.models import PenugasanPetugas

        output = StringIO()
        call_command("seed_data", stdout=output)

        self.assertEqual(Posyandu.objects.count(), 12)
        self.assertEqual(Petugas.objects.filter(level="bidan").count(), 2)
        self.assertEqual(Petugas.objects.filter(level="kader").count(), 12)
        self.assertEqual(Peserta.objects.count(), 72)
        self.assertEqual(Peserta.objects.filter(status_peserta="balita").count(), 48)
        self.assertEqual(Peserta.objects.filter(status_peserta="bumil").count(), 24)
        self.assertEqual(JadwalKegiatan.objects.count(), 24)
        self.assertEqual(WHOStandardPBU.objects.count(), 3714)

        coverage_counts = sorted(
            PenugasanPetugas.objects.filter(
                petugas__level="bidan",
                
            )
            .values("petugas_id")
            .annotate(total=Count("posyandu_id"))
            .values_list("total", flat=True)
        )
        self.assertEqual(coverage_counts, [6, 6])
        self.assertIn("SEED DATA SELESAI", output.getvalue())

        # Jalankan ulang: tidak menggandakan data demo.
        call_command("seed_data", stdout=StringIO())
        self.assertEqual(Posyandu.objects.count(), 12)
        self.assertEqual(Peserta.objects.count(), 72)
        self.assertEqual(JadwalKegiatan.objects.count(), 24)


class MasterDataExportTests(TestCase):
    def setUp(self):
        from apps.posyandu.services import set_bidan_coverage

        self.admin = User.objects.create_user("admin-export-master", password="aman-12345")
        Petugas.objects.create(user=self.admin, nama="Admin Export", level="admin")
        self.pos1 = Posyandu.objects.create(nama="Pos Export 1", desa="Kelurahan Uji", alamat="Alamat 1")
        self.pos2 = Posyandu.objects.create(nama="Pos Export 2", desa="Kelurahan Uji", alamat="Alamat 2")
        Petugas.objects.create(
            nama="Kader Export",
            level="kader",
            posyandu=self.pos1,
            no_telp="081200000001",
        )
        self.bidan = Petugas.objects.create(
            nama="Bidan Export",
            level="bidan",
            posyandu=self.pos1,
            no_telp="081200000002",
        )
        set_bidan_coverage(self.bidan, [self.pos1.pk, self.pos2.pk])
        self.client.force_login(self.admin)

    def test_export_excel_master_kader_bidan_dan_posyandu(self):
        from io import BytesIO
        from openpyxl import load_workbook

        urls = [
            reverse("posyandu:export_petugas_excel", args=["kader"]),
            reverse("posyandu:export_petugas_excel", args=["bidan"]),
            reverse("posyandu:export_posyandu_excel"),
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertIn("spreadsheetml", response["Content-Type"])
                self.assertIn("attachment", response["Content-Disposition"])
                workbook = load_workbook(BytesIO(response.content))
                self.assertGreater(workbook.active.max_row, 4)

    def test_export_pdf_master_kader_bidan_dan_posyandu(self):
        urls = [
            reverse("posyandu:export_petugas_pdf", args=["kader"]),
            reverse("posyandu:export_petugas_pdf", args=["bidan"]),
            reverse("posyandu:export_posyandu_pdf"),
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/pdf")
                self.assertTrue(response.content.startswith(b"%PDF"))

    def test_export_master_mengikuti_filter_search(self):
        response = self.client.get(
            reverse("posyandu:export_posyandu_excel"),
            {"search": "Export 1"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertGreater(len(response.content), 0)



class LaporanTriwulanPerPosyanduTests(TestCase):
    def setUp(self):
        self.pos_a = Posyandu.objects.create(nama="Posyandu Anggrek")
        self.pos_b = Posyandu.objects.create(nama="Posyandu Melati")
        self.user = User.objects.create_user(
            "admin-triwulan-per-pos",
            password="aman-12345",
        )
        Petugas.objects.create(
            user=self.user,
            nama="Admin Triwulan",
            level="admin",
        )
        self.client.force_login(self.user)

    def test_admin_melihat_satu_posyandu_per_halaman_laporan(self):
        response = self.client.get(
            reverse("posyandu:laporan_triwulan"),
            {"tahun": date.today().year, "triwulan": 1, "pos_page": 1},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["posyandu_page"].paginator.count, 2)
        self.assertEqual(response.context["selected_posyandu"], self.pos_a)
        self.assertEqual(response.context["posyandu_nama"], self.pos_a.nama)
        self.assertNotIn(self.pos_b.nama, response.context["posyandu_nama"])

    def test_halaman_posyandu_berikutnya_mengganti_cakupan_laporan(self):
        response = self.client.get(
            reverse("posyandu:laporan_triwulan"),
            {"tahun": date.today().year, "triwulan": 1, "pos_page": 2},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_posyandu"], self.pos_b)
        self.assertEqual(response.context["posyandu_nama"], self.pos_b.nama)


class LaporanBalitaAdminFilterTests(TestCase):
    def setUp(self):
        from apps.peserta.models import Peserta
        from apps.pemeriksaan.models import PemeriksaanBalita

        self.PemeriksaanBalita = PemeriksaanBalita
        self.pos_a = Posyandu.objects.create(nama="Pos Filter A", desa="Kelurahan Uji")
        self.pos_b = Posyandu.objects.create(nama="Pos Filter B", desa="Kelurahan Uji")

        self.user = User.objects.create_user("admin-filter-balita", password="aman-12345")
        self.petugas = Petugas.objects.create(
            user=self.user, nama="Admin Filter Balita", level="admin"
        )
        self.client.force_login(self.user)

        self.peserta_a = Peserta.objects.create(
            posko=self.pos_a,
            nama_peserta="Balita Filter A",
            tgl_lahir=date(2024, 1, 10),
            alamat="Alamat A",
            no_hp="081200000101",
            status_peserta="balita",
            jenis_kelamin="Laki-Laki",
        )
        self.peserta_b = Peserta.objects.create(
            posko=self.pos_b,
            nama_peserta="Balita Filter B",
            tgl_lahir=date(2024, 2, 10),
            alamat="Alamat B",
            no_hp="081200000102",
            status_peserta="balita",
            jenis_kelamin="Perempuan",
        )

        self.jadwal_a = JadwalKegiatan.objects.create(
            tgl_kegiatan=date(2026, 9, 5),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_a,
        )
        self.jadwal_b = JadwalKegiatan.objects.create(
            tgl_kegiatan=date(2026, 9, 6),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_b,
        )

        tz = timezone.get_current_timezone()
        PemeriksaanBalita.objects.create(
            peserta=self.peserta_a,
            petugas=self.petugas,
            jadwal=self.jadwal_a,
            tgl_pemeriksaan=timezone.make_aware(datetime(2026, 9, 5, 9, 0), tz),
            berat_badan=10.0,
            tinggi_badan=80.0,
            status_gizi="normal",
        )
        PemeriksaanBalita.objects.create(
            peserta=self.peserta_b,
            petugas=self.petugas,
            jadwal=self.jadwal_b,
            tgl_pemeriksaan=timezone.make_aware(datetime(2026, 9, 6, 9, 0), tz),
            berat_badan=9.0,
            tinggi_badan=77.0,
            status_gizi="stunting",
        )

    def test_filter_balita_hanya_bulan_tahun_dan_posyandu(self):
        response = self.client.get(
            reverse("posyandu:laporan_balita"),
            {"bulan": 9, "tahun": 2026, "posyandu": self.pos_a.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_data"], 1)
        self.assertEqual(response.context["selected_posyandu_id"], self.pos_a.pk)
        self.assertContains(response, "Balita Filter A")
        self.assertNotContains(response, "Balita Filter B")
        self.assertContains(response, 'name="bulan"')
        self.assertContains(response, 'name="tahun"')
        self.assertContains(response, 'name="posyandu"')
        self.assertNotContains(response, 'name="tanggal_awal"')
        self.assertNotContains(response, 'name="tanggal_akhir"')

    def test_export_excel_balita_mengikuti_filter_posyandu(self):
        response = self.client.get(
            reverse("posyandu:export_balita_excel"),
            {"bulan": 9, "tahun": 2026, "posyandu": self.pos_a.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])

    def test_tema_laporan_imunisasi_admin_mengikuti_base(self):
        response = self.client.get(reverse("posyandu:laporan_imunisasi"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "bg-primary")
        self.assertNotContains(response, "min-h-screen bg-background-light")
        self.assertNotContains(response, "bg-rose-600")


class LaporanImunisasiAdminFilterTests(TestCase):
    def setUp(self):
        self.pos_a = Posyandu.objects.create(nama="Pos Imunisasi A", desa="Kelurahan Uji")
        self.pos_b = Posyandu.objects.create(nama="Pos Imunisasi B", desa="Kelurahan Uji")
        self.user = User.objects.create_user("admin-filter-imunisasi", password="aman-12345")
        Petugas.objects.create(
            user=self.user,
            nama="Admin Filter Imunisasi",
            level="admin",
        )
        self.client.force_login(self.user)

        self.jadwal_a = JadwalKegiatan.objects.create(
            tgl_kegiatan=date(2026, 9, 5),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Imunisasi",
            posyandu=self.pos_a,
        )
        self.jadwal_b = JadwalKegiatan.objects.create(
            tgl_kegiatan=date(2026, 9, 6),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Imunisasi",
            posyandu=self.pos_b,
        )
        JadwalKegiatan.objects.create(
            tgl_kegiatan=date(2026, 8, 20),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Imunisasi",
            posyandu=self.pos_a,
        )

    def test_filter_imunisasi_hanya_bulan_tahun_dan_posyandu(self):
        response = self.client.get(
            reverse("posyandu:laporan_imunisasi"),
            {"bulan": 9, "tahun": 2026, "posyandu": self.pos_a.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_posyandu_id"], self.pos_a.pk)
        self.assertEqual(response.context["total_jadwal"], 1)
        self.assertContains(response, 'name="bulan"')
        self.assertContains(response, 'name="tahun"')
        self.assertContains(response, 'name="posyandu"')
        self.assertNotContains(response, 'name="tanggal_awal"')
        self.assertNotContains(response, 'name="tanggal_akhir"')
        self.assertContains(response, "Semua Posyandu · 1 PDF")
        self.assertContains(response, "scope=all")
        self.assertContains(response, "scope=selected")

    def test_pdf_imunisasi_posyandu_terpilih_hanya_satu_dokumen(self):
        from io import BytesIO
        from pypdf import PdfReader

        response = self.client.get(
            reverse("posyandu:cetak_imunisasi_pdf"),
            {
                "scope": "selected",
                "status_peserta": "balita",
                "bulan": 9,
                "tahun": 2026,
                "posyandu": self.pos_a.pk,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/pdf", response["Content-Type"])
        reader = PdfReader(BytesIO(response.content))
        self.assertEqual(len(reader.pages), 1)

    def test_pdf_imunisasi_semua_posyandu_satu_file_per_posyandu(self):
        from io import BytesIO
        from pypdf import PdfReader

        response = self.client.get(
            reverse("posyandu:cetak_imunisasi_pdf"),
            {
                "scope": "all",
                "status_peserta": "balita",
                "bulan": 9,
                "tahun": 2026,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/pdf", response["Content-Type"])
        self.assertIn("semua_posyandu", response["Content-Disposition"])
        reader = PdfReader(BytesIO(response.content))
        self.assertEqual(len(reader.pages), 2)


class JadwalSinkronisasiAdminKaderTests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.pos_a = Posyandu.objects.create(nama="Pos Sinkron A", desa="Kelurahan Uji")
        self.pos_b = Posyandu.objects.create(nama="Pos Sinkron B", desa="Kelurahan Uji")

        self.admin_user = User.objects.create_user("admin-sinkron", password="aman-12345")
        Petugas.objects.create(
            user=self.admin_user,
            nama="Admin Sinkron",
            level="admin",
        )

        self.kader_user = User.objects.create_user("kader-sinkron", password="aman-12345")
        self.kader = Petugas.objects.create(
            user=self.kader_user,
            nama="Kader Sinkron",
            level="kader",
            posyandu=self.pos_a,
        )

        self.jadwal_dekat = JadwalKegiatan.objects.create(
            tgl_kegiatan=self.today + timedelta(days=3),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_a,
        )
        self.jadwal_10_oktober_simulasi = JadwalKegiatan.objects.create(
            tgl_kegiatan=self.today + timedelta(days=18),
            jam_mulai=time(9),
            jam_selesai=time(11),
            jns_kegiatan="Imunisasi",
            posyandu=self.pos_a,
        )
        self.jadwal_pos_lain = JadwalKegiatan.objects.create(
            tgl_kegiatan=self.today + timedelta(days=5),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_b,
        )
        self.jadwal_lama = JadwalKegiatan.objects.create(
            tgl_kegiatan=self.today - timedelta(days=2),
            jam_mulai=time(8),
            jam_selesai=time(10),
            jns_kegiatan="Pemeriksaan Rutin",
            posyandu=self.pos_a,
        )

    def test_jadwal_posyandu_yang_sama_muncul_di_admin_dan_kader(self):
        self.client.force_login(self.admin_user)
        admin_response = self.client.get(
            reverse("posyandu:jadwal_posyandu_admin"),
            {"status": "mendatang", "posyandu": self.pos_a.pk},
        )
        admin_ids = [item.pk for item in admin_response.context["page_obj"].object_list]
        self.assertIn(self.jadwal_dekat.pk, admin_ids)
        self.assertIn(self.jadwal_10_oktober_simulasi.pk, admin_ids)
        self.assertNotIn(self.jadwal_lama.pk, admin_ids)

        self.client.force_login(self.kader_user)
        kader_response = self.client.get(reverse("posyandu:list_jadwal"))
        kader_ids = [item.pk for item in kader_response.context["page_obj"].object_list]
        self.assertIn(self.jadwal_dekat.pk, kader_ids)
        self.assertIn(self.jadwal_10_oktober_simulasi.pk, kader_ids)
        self.assertNotIn(self.jadwal_pos_lain.pk, kader_ids)
        self.assertNotIn(self.jadwal_lama.pk, kader_ids)

    def test_admin_dan_kader_memakai_urutan_mendatang_yang_sama(self):
        self.client.force_login(self.admin_user)
        admin_response = self.client.get(
            reverse("posyandu:jadwal_posyandu_admin"),
            {"status": "mendatang", "posyandu": self.pos_a.pk},
        )
        admin_ids = [item.pk for item in admin_response.context["page_obj"].object_list]

        self.client.force_login(self.kader_user)
        kader_response = self.client.get(reverse("posyandu:list_jadwal"))
        kader_ids = [item.pk for item in kader_response.context["page_obj"].object_list]

        self.assertEqual(admin_ids, kader_ids)
        self.assertEqual(admin_ids[0], self.jadwal_dekat.pk)
        self.assertEqual(admin_ids[1], self.jadwal_10_oktober_simulasi.pk)

    def test_admin_menampilkan_nama_kader_penerima_jadwal(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(
            reverse("posyandu:jadwal_posyandu_admin"),
            {"status": "mendatang", "posyandu": self.pos_a.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Kader Sinkron")

    def test_admin_default_hanya_menampilkan_jadwal_mendatang(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(
            reverse("posyandu:jadwal_posyandu_admin"),
            {"posyandu": self.pos_a.pk},
        )
        ids = [item.pk for item in response.context["page_obj"].object_list]
        self.assertEqual(response.context["status_filter"], "mendatang")
        self.assertNotIn(self.jadwal_lama.pk, ids)


class KaderSingleOperationalPosyanduTests(TestCase):
    def test_posyandu_utama_kader_menjadi_sumber_tunggal_operasional(self):
        from apps.posyandu.models import PenugasanPetugas
        from common.access import active_posyandu_ids

        utama = Posyandu.objects.create(nama="Pos Operasional Utama")
        stale = Posyandu.objects.create(nama="Pos Penugasan Lama")
        user = User.objects.create_user("kader-operasional", password="aman-12345")
        petugas = Petugas.objects.create(
            user=user,
            nama="Kader Operasional",
            level="kader",
            posyandu=utama,
        )
        # Simulasikan data legacy yang masih memiliki penugasan tambahan kedua.
        PenugasanPetugas.objects.create(
            petugas=petugas,
            posyandu=stale,
            tanggal_daftar=timezone.localdate(),
        )

        self.assertEqual(active_posyandu_ids(petugas), {utama.pk})
