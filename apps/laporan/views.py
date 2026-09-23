import calendar
from datetime import date

from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Count, Prefetch
import io
from django.core.paginator import Paginator
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from xhtml2pdf import pisa
from io import BytesIO
from pypdf import PdfReader, PdfWriter

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils.text import slugify
from django.contrib.auth.decorators import login_required
from django.urls import reverse

from common.decorators import petugas_required
from common.access import report_posyandu_ids

from apps.posyandu.models import JadwalKegiatan, Posyandu
from apps.peserta.models import Peserta
from apps.pemeriksaan.models import Imunisasi, PemeriksaanBalita, PemeriksaanBumil
from apps.pemeriksaan.helper.constant import JADWAL_VAKSIN_BALITA

from .helper.utils import build_context_laporan_triwulan, _ambil_parameter
from .helper.kms_standards import bbu_series, pbu_tbu_series, imtu_series


# =============================================
# HELPER
# =============================================
def get_petugas_or_403(request):
    return getattr(request.user, "petugas", None)


def get_posyandu_ids(petugas):
    return report_posyandu_ids(petugas)


def _filter_periode(request):
    today = date.today()
    try:
        bulan = int(request.GET.get("bulan", today.month))
    except (TypeError, ValueError):
        bulan = today.month
    try:
        tahun = int(request.GET.get("tahun", today.year))
    except (TypeError, ValueError):
        tahun = today.year
    return min(max(bulan, 1), 12), min(max(tahun, 2000), today.year + 1)


# =============================================
# DASHBOARD LAPORAN
# =============================================
@login_required
@petugas_required
def laporan_index(request):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)

    # Jadwal bulan ini
    today = date.today()
    jadwal_list = JadwalKegiatan.objects.filter(
        posyandu_id__in=posyandu_ids,
        tgl_kegiatan__year=today.year,
        tgl_kegiatan__month=today.month,
    ).select_related("posyandu").order_by("-tgl_kegiatan")

    # Ringkasan bulan ini
    total_balita = Peserta.objects.filter(
        posko_id__in=posyandu_ids,
        status_peserta="balita"
    ).count()

    total_bumil = Peserta.objects.filter(
        posko_id__in=posyandu_ids,
        status_peserta="bumil"
    ).count()

    total_imunisasi = Imunisasi.objects.filter(
        jadwal__posyandu_id__in=posyandu_ids,
        tgl_pemberian__year=today.year,
        tgl_pemberian__month=today.month,
    ).count()

    total_pemeriksaan_balita = PemeriksaanBalita.objects.filter(
        jadwal__posyandu_id__in=posyandu_ids,
        tgl_pemeriksaan__year=today.year,
        tgl_pemeriksaan__month=today.month,
    ).count()

    return render(request, "laporan/laporan.html", {
        "jadwal_list": jadwal_list,
        "total_balita": total_balita,
        "total_bumil": total_bumil,
        "total_imunisasi": total_imunisasi,
        "total_pemeriksaan_balita": total_pemeriksaan_balita,
        "bulan": today.strftime("%B %Y"),
    })


# =============================================
# REKAPITULASI PER JADWAL
# =============================================
@login_required
@petugas_required
def laporan_jadwal(request, id_jadwal):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)

    jadwal = get_object_or_404(
        JadwalKegiatan,
        id_jadwal=id_jadwal,
        posyandu_id__in=posyandu_ids
    )

    status_peserta = request.GET.get("status_peserta", "balita")
    if status_peserta not in {"balita", "bumil"}:
        status_peserta = "balita"

    # Peserta terdaftar
    peserta_list = Peserta.objects.filter(
        posko=jadwal.posyandu,
        status_peserta=status_peserta
    ).order_by("nama_peserta")

    total_terdaftar = peserta_list.count()

    if jadwal.jns_kegiatan == "Imunisasi":
        
        # Rekap imunisasi
        imunisasi_qs = Imunisasi.objects.filter(
            jadwal=jadwal,
            peserta__status_peserta=status_peserta
        ).select_related("peserta")

        hadir_ids = set(imunisasi_qs.values_list("peserta_id", flat=True))
        total_hadir = len(hadir_ids)

        # Map peserta -> vaksin
        imunisasi_map = {}
        for imun in imunisasi_qs:
            if imun.peserta_id not in imunisasi_map:
                imunisasi_map[imun.peserta_id] = []
            imunisasi_map[imun.peserta_id].append(imun.get_jenis_vaksin_display())

        for p in peserta_list:
            p.hadir = p.id in hadir_ids
            p.vaksin_list = imunisasi_map.get(p.id, [])

        rekap_vaksin = imunisasi_qs.values(
            "jenis_vaksin"
        ).annotate(
            jumlah=Count("id_imunisasi")
        ).order_by("-jumlah")

    else:
        # Rekap pemeriksaan rutin
        PemeriksaanModel = {
            "balita": PemeriksaanBalita,
            "bumil": PemeriksaanBumil,
        }.get(status_peserta)

        pemeriksaan_qs = PemeriksaanModel.objects.filter(
            jadwal=jadwal
        ).select_related("peserta")

        hadir_ids = set(pemeriksaan_qs.values_list("peserta_id", flat=True))
        total_hadir = len(hadir_ids)

        periksa_map = {p.peserta_id: p for p in pemeriksaan_qs}

        for p in peserta_list:
            p.hadir = p.id in hadir_ids
            p.pemeriksaan = periksa_map.get(p.id)

        rekap_vaksin = None

    return render(request, "laporan/laporan_jadwal.html", {
        "jadwal": jadwal,
        "peserta_list": peserta_list,
        "status_peserta": status_peserta,
        "total_terdaftar": total_terdaftar,
        "total_hadir": total_hadir,
        "total_tidak_hadir": total_terdaftar - total_hadir,
        "persen_hadir": round((total_hadir / total_terdaftar * 100) if total_terdaftar else 0),
        "rekap_vaksin": rekap_vaksin,
    })


# =============================================
# LAPORAN BULANAN
# =============================================
@login_required
@petugas_required
def laporan_bulanan(request):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)

    # Filter bulan & tahun
    today = date.today()
    bulan, tahun = _filter_periode(request)

    jadwal_list = JadwalKegiatan.objects.filter(
        posyandu_id__in=posyandu_ids,
        tgl_kegiatan__year=tahun,
        tgl_kegiatan__month=bulan,
    ).select_related("posyandu").order_by("tgl_kegiatan")

    # Rekap per jadwal
    rekap = []
    for jadwal in jadwal_list:
        if jadwal.jns_kegiatan == "Imunisasi":
            hadir = Imunisasi.objects.filter(
                jadwal=jadwal
            ).values("peserta_id").distinct().count()
        else:
            hadir_balita = PemeriksaanBalita.objects.filter(jadwal=jadwal).count()
            hadir_bumil  = PemeriksaanBumil.objects.filter(jadwal=jadwal).count()
            hadir = hadir_balita + hadir_bumil

        terdaftar = Peserta.objects.filter(
            posko=jadwal.posyandu
        ).count()

        rekap.append({
            "jadwal": jadwal,
            "terdaftar": terdaftar,
            "hadir": hadir,
            "tidak_hadir": terdaftar - hadir,
            "persen": round((hadir / terdaftar * 100) if terdaftar else 0),
        })

    # Pilihan bulan untuk filter
    bulan_choices = [(i, calendar.month_name[i]) for i in range(1, 13)]
    tahun_choices = list(range(today.year - 2, today.year + 1))

    return render(request, "laporan/laporan_bulanan2.html", {
        "rekap": rekap,
        "bulan": bulan,
        "tahun": tahun,
        "bulan_choices": bulan_choices,
        "tahun_choices": tahun_choices,
        "bulan_label": calendar.month_name[bulan],
    })


# =============================================
# RIWAYAT IMUNISASI PER PESERTA
# =============================================
# Grup vaksin program bayi/baduta untuk kolom tabel rekap per jadwal.
# Nama harus sama persis dengan pilihan pada model Imunisasi agar input Kader
# selalu terbaca oleh laporan.
VAKSIN_BALITA_GROUPS = {
    "hb0":       ["HB-0"],
    "bcg":       ["BCG"],
    "dpt":       ["DPT-HB-Hib 1", "DPT-HB-Hib 2", "DPT-HB-Hib 3", "DPT-HB-Hib 4"],
    "polio":     ["bOPV 1", "bOPV 2", "bOPV 3", "bOPV 4", "IPV 1", "IPV 2"],
    "mr":        ["MR 1", "MR 2"],
    "pcv":       ["PCV 1", "PCV 2", "PCV 3"],
    "rotavirus": ["Rotavirus 1", "Rotavirus 2", "Rotavirus 3"],
}

# Urutan kolom ringkasan riwayat vaksin per sasaran.
KELENGKAPAN_GROUPS = {
    "HB-0":          ["HB-0"],
    "BCG":           ["BCG"],
    "DPT-HB-Hib":    ["DPT-HB-Hib 1", "DPT-HB-Hib 2", "DPT-HB-Hib 3", "DPT-HB-Hib 4"],
    "Polio":         ["bOPV 1", "bOPV 2", "bOPV 3", "bOPV 4", "IPV 1", "IPV 2"],
    "MR":            ["MR 1", "MR 2"],
    "PCV":           ["PCV 1", "PCV 2", "PCV 3"],
    "Rotavirus":     ["Rotavirus 1", "Rotavirus 2", "Rotavirus 3"],
}

# Pilihan vaksin untuk filter laporan. JE tetap tersedia untuk pencatatan manual
# karena penggunaannya bergantung program/wilayah. TT dipertahankan untuk bumil.
VAKSIN_FILTER_CHOICES = [
    ("HB-0", "HB-0"),
    ("BCG", "BCG"),
    ("bOPV 1", "bOPV 1"),
    ("bOPV 2", "bOPV 2"),
    ("bOPV 3", "bOPV 3"),
    ("bOPV 4", "bOPV 4"),
    ("IPV 1", "IPV 1"),
    ("IPV 2", "IPV 2"),
    ("DPT-HB-Hib 1", "DPT-HB-Hib 1"),
    ("DPT-HB-Hib 2", "DPT-HB-Hib 2"),
    ("DPT-HB-Hib 3", "DPT-HB-Hib 3"),
    ("DPT-HB-Hib 4", "DPT-HB-Hib 4"),
    ("PCV 1", "PCV 1"),
    ("PCV 2", "PCV 2"),
    ("PCV 3", "PCV 3"),
    ("Rotavirus 1", "Rotavirus 1"),
    ("Rotavirus 2", "Rotavirus 2"),
    ("Rotavirus 3", "Rotavirus 3"),
    ("MR 1", "MR 1"),
    ("MR 2", "MR 2"),
    ("JE", "JE (wilayah/program tertentu)"),
    ("TT-1", "TT-1 (Tetanus Toksoid)"),
    ("TT-2", "TT-2 (Tetanus Toksoid)"),
    ("TT-3", "TT-3 (Tetanus Toksoid)"),
    ("TT-4", "TT-4 (Tetanus Toksoid)"),
    ("TT-5", "TT-5 (Tetanus Toksoid)"),
]


def _usia_bulan_pada(tgl_lahir, tanggal_acuan):
    """Hitung umur bulan penuh pada tanggal acuan laporan."""
    if not tgl_lahir or tanggal_acuan < tgl_lahir:
        return None
    bulan = (tanggal_acuan.year - tgl_lahir.year) * 12 + (tanggal_acuan.month - tgl_lahir.month)
    if tanggal_acuan.day < tgl_lahir.day:
        bulan -= 1
    return max(bulan, 0)


def _vaksin_wajib_sesuai_usia(tgl_lahir, tanggal_acuan):
    """Vaksin program yang sudah jatuh tempo sampai tanggal acuan.

    JE tidak dimasukkan karena bersifat wilayah/program tertentu.
    """
    usia_bulan = _usia_bulan_pada(tgl_lahir, tanggal_acuan)
    if usia_bulan is None:
        return set()

    wajib = set(JADWAL_VAKSIN_BALITA.get("0-24jam", []))
    for usia, vaksin_list in JADWAL_VAKSIN_BALITA.items():
        if isinstance(usia, int) and usia <= usia_bulan:
            wajib.update(vaksin_list)
    return wajib

# =============================================
# LAPORAN IMUNISASI
# =============================================
@login_required
@petugas_required
def laporan_imunisasi_bulanan(request):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    # Admin Kelurahan memiliki halaman laporan imunisasi sendiri.
    # Route laporan umum ini dipertahankan untuk Kader/kompatibilitas lama.
    if petugas.level == "admin":
        query = request.GET.urlencode()
        target = reverse("posyandu:laporan_imunisasi")
        if query:
            target = f"{target}?{query}"
        return redirect(target)
 
    posyandu_ids = get_posyandu_ids(petugas)
 
    # --- Filter bulan, tahun, jenis vaksin ---
    today = date.today()
    bulan, tahun = _filter_periode(request)
    vaksin = request.GET.get("vaksin", "")  # "" = semua vaksin
 
    # Jadwal Imunisasi dalam periode
    jadwal_qs = JadwalKegiatan.objects.filter(
        posyandu_id__in=posyandu_ids,
        tgl_kegiatan__year=tahun,
        tgl_kegiatan__month=bulan,
        jns_kegiatan="Imunisasi",
    ).select_related("posyandu").order_by("tgl_kegiatan")
 
    # -----------------------------------------------
    # REKAP PER JADWAL (tabel atas)
    # -----------------------------------------------
    rekap = []
    for jadwal in jadwal_qs:
        imunisasi_qs = Imunisasi.objects.filter(jadwal=jadwal)
        if vaksin:
            imunisasi_qs = imunisasi_qs.filter(jenis_vaksin=vaksin)
 
        sasaran = Peserta.objects.filter(posko=jadwal.posyandu).count()
 
        # Hitung dosis per grup vaksin
        def hitung_grup(keys):
            return imunisasi_qs.filter(jenis_vaksin__in=keys).count()
 
        hb0       = hitung_grup(VAKSIN_BALITA_GROUPS["hb0"])
        bcg       = hitung_grup(VAKSIN_BALITA_GROUPS["bcg"])
        dpt       = hitung_grup(VAKSIN_BALITA_GROUPS["dpt"])
        polio     = hitung_grup(VAKSIN_BALITA_GROUPS["polio"])
        mr        = hitung_grup(VAKSIN_BALITA_GROUPS["mr"])
        pcv       = hitung_grup(VAKSIN_BALITA_GROUPS["pcv"])
        rotavirus = hitung_grup(VAKSIN_BALITA_GROUPS["rotavirus"])
 
        total_dosis = imunisasi_qs.count()
 
        # Cakupan: peserta unik yang hadir vs sasaran terdaftar
        peserta_hadir = imunisasi_qs.values("peserta_id").distinct().count()
        persen = round((peserta_hadir / sasaran * 100) if sasaran else 0)
 
        rekap.append({
            "jadwal":      jadwal,
            "sasaran":     sasaran,
            "hb0":         hb0,
            "bcg":         bcg,
            "dpt":         dpt,
            "polio":       polio,
            "mr":          mr,
            "pcv":         pcv,
            "rotavirus":   rotavirus,
            "total_dosis": total_dosis,
            "persen":      persen,
        })
 
    # -----------------------------------------------
    # STATUS KELENGKAPAN PER SASARAN (tabel bawah)
    # -----------------------------------------------
    # Semua peserta balita yang sudah lahir pada akhir periode laporan.
    batas_akhir = date(tahun, bulan, calendar.monthrange(tahun, bulan)[1])
    peserta_balita = Peserta.objects.filter(
        posko_id__in=posyandu_ids,
        status_peserta="balita",
        tgl_lahir__lte=batas_akhir,
    ).select_related("posko").order_by("nama_peserta")
 
    # Semua imunisasi peserta tersebut s.d. akhir bulan terpilih.
    imunisasi_all = Imunisasi.objects.filter(
        peserta__in=peserta_balita,
        tgl_pemberian__lte=batas_akhir,
    ).values("peserta_id", "jenis_vaksin")
 
    # Buat lookup: peserta_id -> set jenis_vaksin yang sudah diterima
    vaksin_per_peserta: dict[int, set] = {}
    for row in imunisasi_all:
        vaksin_per_peserta.setdefault(row["peserta_id"], set()).add(row["jenis_vaksin"])
 
    sasaran_list = []
    for p in peserta_balita:
        sudah = vaksin_per_peserta.get(p.pk, set())
 
        # Status per grup kelengkapan
        status_vaksin = []
        for grup_keys in KELENGKAPAN_GROUPS.values():
            sudah_dalam_grup = [v for v in grup_keys if v in sudah]
            if len(sudah_dalam_grup) == len(grup_keys):
                status_vaksin.append("lengkap")
            elif sudah_dalam_grup:
                status_vaksin.append("sebagian")
            else:
                status_vaksin.append("belum")
 
        # Status akhir dinilai berdasarkan vaksin yang sudah jatuh tempo menurut
        # usia anak pada akhir periode laporan, bukan memaksa bayi muda menerima
        # seluruh dosis sampai usia 18 bulan.
        usia_bulan = _usia_bulan_pada(p.tgl_lahir, batas_akhir)
        vaksin_wajib = _vaksin_wajib_sesuai_usia(p.tgl_lahir, batas_akhir)
        status_lengkap = bool(vaksin_wajib) and vaksin_wajib.issubset(sudah)
 
        sasaran_list.append({
            "nama_anak":      p.nama_peserta,
            "tgl_lahir":      p.tgl_lahir,
            "usia_bulan":     usia_bulan,
            "posyandu":       p.posko.nama if p.posko else "-",
            "status_vaksin":  status_vaksin,
            "status_lengkap": status_lengkap,
            "jumlah_wajib":   len(vaksin_wajib),
            "jumlah_terpenuhi": len(vaksin_wajib.intersection(sudah)),
        })
 
    # -----------------------------------------------
    # STATISTIK RINGKASAN (kartu atas)
    # -----------------------------------------------
    total_sasaran    = peserta_balita.count()
    sudah_imunisasi  = sum(1 for s in sasaran_list if any(v != "belum" for v in s["status_vaksin"]))
    belum_imunisasi  = total_sasaran - sudah_imunisasi
    lengkap_count    = sum(1 for s in sasaran_list if s["status_lengkap"])
    persen_sesuai_usia = round((lengkap_count / total_sasaran * 100) if total_sasaran else 0)
 
    # -----------------------------------------------
    # PILIHAN FILTER
    # -----------------------------------------------
    bulan_choices = [(i, calendar.month_name[i]) for i in range(1, 13)]
    tahun_choices = list(range(today.year - 2, today.year + 1))
 
    return render(request, "laporan/laporan_imunisasi.html", {
        # Filter aktif
        "bulan":        bulan,
        "tahun":        tahun,
        "vaksin":       vaksin,
        "bulan_label":  calendar.month_name[bulan],
        # Pilihan dropdown
        "bulan_choices":  bulan_choices,
        "tahun_choices":  tahun_choices,
        "vaksin_choices": VAKSIN_FILTER_CHOICES,
        # Data tabel
        "rekap":        rekap,
        "sasaran_list": sasaran_list,
        # Kartu statistik
        "stats": {
            "total_sasaran":   total_sasaran,
            "sudah_imunisasi": sudah_imunisasi,
            "belum_imunisasi": belum_imunisasi,
            "persen_sesuai_usia": persen_sesuai_usia,
        },
        # Label kolom kelengkapan (urutan sama dengan status_vaksin)
        "kelengkapan_headers": list(KELENGKAPAN_GROUPS.keys()),
    })


# =============================================
# KMS
# =============================================
@login_required
@petugas_required
def laporan_kms(request, id_peserta):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)
    peserta = get_object_or_404(
        Peserta,
        id=id_peserta,
        status_peserta="balita",
        posko_id__in=posyandu_ids,
    )

    riwayat = list(
        PemeriksaanBalita.objects.filter(peserta=peserta)
        .order_by("usia_bulan", "tgl_pemeriksaan")
    )

    gender_text = str(peserta.jenis_kelamin or "").strip().lower()
    gender_code = "P" if gender_text.startswith("p") else "L"

    # Data pemeriksaan dipertahankan berdasarkan usia bulan penuh yang sudah
    # disimpan pada PemeriksaanBalita. Titik di luar 0-60 bulan tidak diplot
    # karena KMS balita mengacu pada standar antropometri 0-60 bulan.
    titik = []
    for row in riwayat:
        if row.usia_bulan is None or not (0 <= row.usia_bulan <= 60):
            continue
        imt = None
        if (
            row.berat_badan is not None
            and row.tinggi_badan is not None
            and float(row.berat_badan) > 0
            and float(row.tinggi_badan) > 0
        ):
            tinggi_meter = float(row.tinggi_badan) / 100.0
            imt = round(float(row.berat_badan) / (tinggi_meter ** 2), 2)

        titik.append({
            "bulan": int(row.usia_bulan),
            "berat": row.berat_badan,
            "tinggi": row.tinggi_badan,
            "imt": imt,
            "tanggal": row.tgl_pemeriksaan.strftime("%d-%m-%Y"),
            "z_score": row.z_score,
            "status": row.get_status_gizi_display() if row.status_gizi else "-",
        })

    chart_data = {
        "gender": gender_code,
        "points": titik,
        "bbu": bbu_series(gender_code),
        "pbu_tbu": pbu_tbu_series(gender_code),
        "imtu": imtu_series(gender_code),
    }

    pemeriksaan_terakhir = riwayat[-1] if riwayat else None

    return render(request, "laporan/laporan_kms.html", {
        "peserta": peserta,
        "riwayat": riwayat,
        "chart_data": chart_data,
        "pemeriksaan_terakhir": pemeriksaan_terakhir,
    })


# ============================================================
# VIEW: halaman laporan triwulan
# ============================================================
@login_required
@petugas_required
def laporan_triwulan(request):
    petugas = get_petugas_or_403(request)
    allowed_ids = list(get_posyandu_ids(petugas))
    tahun, triwulan = _ambil_parameter(request)

    daftar_posyandu = list(
        Posyandu.objects
        .filter(pk__in=allowed_ids)
        .order_by("nama")
    )
    is_admin_kelurahan = bool(petugas and petugas.level == "admin")

    # Admin Kelurahan melihat satu Posyandu pada satu waktu agar laporan
    # tidak menumpuk. Kader tetap memakai seluruh Posyandu yang memang
    # menjadi cakupan aksesnya (umumnya satu Posyandu).
    selected_posyandu = None
    context_posyandu_ids = allowed_ids

    if is_admin_kelurahan and daftar_posyandu:
        selected_raw = request.GET.get("posyandu", "").strip()
        try:
            selected_id = int(selected_raw) if selected_raw else None
        except (TypeError, ValueError):
            selected_id = None

        selected_posyandu = next(
            (item for item in daftar_posyandu if item.pk == selected_id),
            None,
        )
        if selected_posyandu is None:
            selected_posyandu = daftar_posyandu[0]

        context_posyandu_ids = [selected_posyandu.pk]

    context = build_context_laporan_triwulan(
        tahun,
        triwulan,
        context_posyandu_ids,
    )
    context.update({
        "is_admin_kelurahan": is_admin_kelurahan,
        "daftar_posyandu_filter": daftar_posyandu,
        "selected_posyandu": selected_posyandu,
        "selected_posyandu_id": selected_posyandu.pk if selected_posyandu else None,
    })

    paginator_balita = Paginator(context['data_balita'], 5)
    context['data_balita'] = paginator_balita.get_page(
        request.GET.get('page_balita', 1)
    )

    paginator_bumil = Paginator(context['data_bumil'], 10)
    context['data_bumil'] = paginator_bumil.get_page(
        request.GET.get('page_bumil', 1)
    )

    return render(request, 'laporan/laporan_triwulan.html', context)



# ============================================================
# VIEW: export PDF
# ============================================================
def _render_pdf_triwulan_satu_posyandu(template, context):
    """Render satu laporan Posyandu menjadi bytes PDF.

    Helper ini sengaja merender satu Posyandu per dokumen sementara agar
    format laporan lama tetap dipertahankan. Dokumen-dokumen sementara
    kemudian digabung sehingga setiap Posyandu selalu dimulai pada halaman
    baru dan metadata header hanya memuat satu nama Posyandu.
    """
    html_string = render_to_string(template, context)
    output = BytesIO()
    status = pisa.CreatePDF(io.StringIO(html_string), dest=output)
    if status.err:
        raise RuntimeError("Gagal membuat PDF laporan triwulan Posyandu.")
    output.seek(0)
    return output


@login_required
@petugas_required
def export_laporan_triwulan_pdf(request):
    petugas = get_petugas_or_403(request)
    allowed_ids = list(get_posyandu_ids(petugas))
    tahun, triwulan = _ambil_parameter(request)
    jenis = request.GET.get("jenis", "balita")
    is_admin_kelurahan = bool(petugas and petugas.level == "admin")

    if jenis == "bumil":
        template = "laporan/pdf_laporan_triwulan_bumil.html"
        jenis_filename = "bumil"
    else:
        template = "laporan/pdf_laporan_triwulan_balita.html"
        jenis_filename = "balita"

    allowed_posyandu = list(
        Posyandu.objects
        .filter(pk__in=allowed_ids)
        .order_by("nama")
    )
    allowed_map = {item.pk: item for item in allowed_posyandu}

    # scope=all hanya dipakai Admin Kelurahan untuk mengunduh seluruh
    # Posyandu sekaligus. Hasil tetap satu PDF, tetapi setiap Posyandu
    # dirender sebagai dokumen terpisah lalu digabung sehingga selalu
    # dimulai pada halaman baru.
    scope = request.GET.get("scope", "selected").strip().lower()
    selected_raw = request.GET.get("posyandu", "").strip()
    try:
        selected_id = int(selected_raw) if selected_raw else None
    except (TypeError, ValueError):
        selected_id = None

    if is_admin_kelurahan and scope == "all":
        export_ids = [item.pk for item in allowed_posyandu]
        filename = (
            f"laporan_triwulan_{jenis_filename}_semua_posyandu_"
            f"{tahun}_TW{triwulan}.pdf"
        )
    elif selected_id is not None and selected_id in allowed_map:
        export_ids = [selected_id]
        pos_slug = slugify(allowed_map[selected_id].nama) or f"posyandu-{selected_id}"
        filename = (
            f"laporan_triwulan_{jenis_filename}_{pos_slug}_"
            f"{tahun}_TW{triwulan}.pdf"
        )
    elif is_admin_kelurahan and allowed_posyandu:
        # Jika Admin membuka URL export tanpa parameter Posyandu,
        # gunakan Posyandu pertama agar tidak kembali menumpuk semua.
        first = allowed_posyandu[0]
        export_ids = [first.pk]
        pos_slug = slugify(first.nama) or f"posyandu-{first.pk}"
        filename = (
            f"laporan_triwulan_{jenis_filename}_{pos_slug}_"
            f"{tahun}_TW{triwulan}.pdf"
        )
    else:
        # Kader tidak diberi filter multi-Posyandu. Export mengikuti
        # cakupan aksesnya seperti sebelumnya.
        export_ids = [item.pk for item in allowed_posyandu]
        filename = f"laporan_triwulan_{jenis_filename}_{tahun}_TW{triwulan}.pdf"

    if not export_ids:
        return HttpResponse(
            "Tidak ada Posyandu yang dapat diekspor.",
            status=404,
            content_type="text/plain; charset=utf-8",
        )

    writer = PdfWriter()
    temporary_buffers = []

    try:
        for posyandu_id in export_ids:
            context = build_context_laporan_triwulan(
                tahun,
                triwulan,
                [posyandu_id],
            )

            pdf_buffer = _render_pdf_triwulan_satu_posyandu(
                template,
                context,
            )
            temporary_buffers.append(pdf_buffer)

            reader = PdfReader(pdf_buffer)
            for page in reader.pages:
                writer.add_page(page)

        merged_pdf = BytesIO()
        writer.write(merged_pdf)
        merged_pdf.seek(0)

    except Exception:
        return HttpResponse(
            "Terjadi kesalahan saat membuat PDF laporan triwulan.",
            status=500,
            content_type="text/plain; charset=utf-8",
        )
    finally:
        for buffer in temporary_buffers:
            buffer.close()

    response = HttpResponse(merged_pdf.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# =============================================
# EXPORT PDF
# =============================================
# Ganti fungsi export PDF
@login_required
@petugas_required
def export_laporan_jadwal_pdf(request, id_jadwal):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)
    jadwal = get_object_or_404( 
        JadwalKegiatan,
        id_jadwal=id_jadwal,
        posyandu_id__in=posyandu_ids
    )

    # 1. Ambil status peserta dari parameter URL (?status_peserta=balita / bumil)
    status_peserta = request.GET.get("status_peserta", "balita")
    if status_peserta not in {"balita", "bumil"}:
        status_peserta = "balita"

    # 2. Query dasar peserta terdaftar di posyandu tersebut
    peserta_queryset = Peserta.objects.filter(
        posko=jadwal.posyandu,
        status_peserta=status_peserta
    ).order_by('nama_peserta')

    # =========================================================
    # KONDISI 1: JIKA JADWAL ADALAH IMUNISASI
    # =========================================================
    if jadwal.jns_kegiatan == "Imunisasi":
        imunisasi_qs = Imunisasi.objects.filter(
            jadwal=jadwal,
            peserta__status_peserta=status_peserta
        ).select_related("peserta")

        hadir_ids = set(imunisasi_qs.values_list("peserta_id", flat=True))

        # Mapping peserta_id -> detail imunisasi
        imunisasi_map = {}
        for imun in imunisasi_qs:
            if imun.peserta_id not in imunisasi_map:
                imunisasi_map[imun.peserta_id] = {
                    'vaksin': [],
                    'usia_saat_vaksin': imun.usia_saat_vaksin
                }
            imunisasi_map[imun.peserta_id]['vaksin'].append(imun.get_jenis_vaksin_display())

        # Pasang data ke objek secara in-memory
        for p in peserta_queryset:
            p.hadir = p.id in hadir_ids
            p_data = imunisasi_map.get(p.id, {'vaksin': [], 'usia_saat_vaksin': None})
            p.vaksin_list = p_data['vaksin']
            p.usia_saat_vaksin = p_data['usia_saat_vaksin']
        
        template_name = "laporan/pdf_imunisasi.html"
        filename_prefix = f"laporan imunisasi {status_peserta} {jadwal.id_jadwal}"

    # =========================================================
    # KONDISI 2: JIKA JADWAL ADALAH PEMERIKSAAN RUTIN
    # =========================================================
    else:
        if status_peserta == 'bumil':
            peserta_queryset = peserta_queryset.prefetch_related(
                Prefetch(
                    'pemeriksaan_bumil', 
                    queryset=PemeriksaanBumil.objects.filter(jadwal=jadwal),
                    to_attr='riwayat_periksa'
                )
            )
        else:
            peserta_queryset = peserta_queryset.prefetch_related(
                Prefetch(
                    'pemeriksaan_balita', 
                    queryset=PemeriksaanBalita.objects.order_by('tgl_pemeriksaan'), 
                    to_attr='riwayat_periksa'
                )
            )
        
        template_name = "laporan/pdf_jadwal.html"
        filename_prefix = f"laporan Bulanan pemeriksaan {status_peserta} {jadwal.id_jadwal}"

    # 3. Render template yang sesuai ke HTML string
    html_string = render_to_string(template_name, {
        "jadwal": jadwal,
        "data_peserta": peserta_queryset,
        "status_peserta": status_peserta,
    })

    buffer = BytesIO()
    pisa.CreatePDF(html_string, dest=buffer)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename_prefix}.pdf"'
    return response


# =============================================
# EXPORT EXCEL
# =============================================
@login_required
@petugas_required
def export_laporan_jadwal_excel(request, id_jadwal):
    petugas = get_petugas_or_403(request)
    if not petugas:
        return render(request, "403.html")

    posyandu_ids = get_posyandu_ids(petugas)
    jadwal = get_object_or_404(
        JadwalKegiatan,
        id_jadwal=id_jadwal,
        posyandu_id__in=posyandu_ids
    )

    status_peserta = request.GET.get("status_peserta", "balita")
    if status_peserta not in {"balita", "bumil"}:
        status_peserta = "balita"

    peserta_list = Peserta.objects.filter(
        posko=jadwal.posyandu,
        status_peserta=status_peserta
    ).order_by("nama_peserta")

    # ✅ Definisikan hadir_ids sesuai jenis kegiatan
    if jadwal.jns_kegiatan == "Imunisasi":
        hadir_ids = set(
            Imunisasi.objects.filter(
                jadwal=jadwal,
                peserta__status_peserta=status_peserta
            ).values_list("peserta_id", flat=True)
        )

        # Map peserta -> vaksin untuk kolom tambahan
        imunisasi_map = {}
        for imun in Imunisasi.objects.filter(jadwal=jadwal, peserta__status_peserta=status_peserta):
            if imun.peserta_id not in imunisasi_map:
                imunisasi_map[imun.peserta_id] = []
            imunisasi_map[imun.peserta_id].append(imun.get_jenis_vaksin_display())

    else:
        PemeriksaanModel = {
            "balita": PemeriksaanBalita,
            "bumil": PemeriksaanBumil,
        }.get(status_peserta)

        hadir_ids = set(
            PemeriksaanModel.objects.filter(
                jadwal=jadwal
            ).values_list("peserta_id", flat=True)
        )
        imunisasi_map = {}

    # Buat workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Kegiatan"

    # Header style
    header_fill = PatternFill("solid", fgColor="16A34A")
    header_font = Font(bold=True, color="FFFFFF")

    # Header kolom — tambah kolom vaksin jika imunisasi
    if jadwal.jns_kegiatan == "Imunisasi":
        headers = ["No", "Nama Peserta", "Usia", "Status", "Keterangan", "Vaksin Diberikan"]
    else:
        headers = ["No", "Nama Peserta", "Usia", "Status", "Keterangan"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Isi data
    for i, p in enumerate(peserta_list, 1):
        hadir = p.id in hadir_ids
        row = [
            i,
            p.nama_peserta,
            p.usia,
            p.status_peserta.title(),
            "Hadir" if hadir else "Tidak Hadir",
        ]

        if jadwal.jns_kegiatan == "Imunisasi":
            vaksin = ", ".join(imunisasi_map.get(p.id, []))
            row.append(vaksin if vaksin else "-")

        ws.append(row)

    # Auto width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="laporan_{jadwal.id_jadwal}.xlsx"'
    wb.save(response)
    return response
